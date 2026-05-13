from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from typing import List
from datetime import datetime
from openpyxl import Workbook, load_workbook
from io import BytesIO
from fastapi.responses import StreamingResponse

from ..database import get_db
from ..models import Book, Reader
from ..schemas import BookCreate

router = APIRouter(prefix="/api/admin", tags=["admin"])


@router.post("/import/books")
async def import_books(file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Only Excel files are supported")
    
    try:
        contents = await file.read()
        workbook = load_workbook(filename=BytesIO(contents))
        sheet = workbook.active
        
        imported_count = 0
        errors = []
        
        headers = [cell.value for cell in sheet[1]]
        header_map = {}
        for idx, header in enumerate(headers):
            if header:
                header_map[str(header).strip().lower()] = idx
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                isbn_idx = header_map.get("isbn")
                title_idx = header_map.get("title") or header_map.get("书名")
                author_idx = header_map.get("author") or header_map.get("作者")
                
                if isbn_idx is None or title_idx is None:
                    continue
                
                isbn = str(row[isbn_idx]) if row[isbn_idx] else None
                title = str(row[title_idx]) if row[title_idx] else None
                
                if not isbn or not title:
                    continue
                
                existing = await db.execute(select(Book).where(Book.isbn == isbn))
                if existing.scalar_one_or_none():
                    continue
                
                publisher_idx = header_map.get("publisher") or header_map.get("出版社")
                category_idx = header_map.get("category") or header_map.get("分类")
                quantity_idx = header_map.get("quantity") or header_map.get("数量")
                
                book = Book(
                    isbn=isbn,
                    title=title,
                    author=str(row[author_idx]) if author_idx and row[author_idx] else "未知作者",
                    publisher=str(row[publisher_idx]) if publisher_idx and row[publisher_idx] else None,
                    category=str(row[category_idx]) if category_idx and row[category_idx] else None,
                    total_quantity=int(row[quantity_idx]) if quantity_idx and row[quantity_idx] else 1,
                    available_quantity=int(row[quantity_idx]) if quantity_idx and row[quantity_idx] else 1,
                )
                
                db.add(book)
                imported_count += 1
                
            except Exception as e:
                errors.append(f"Row {row_idx}: {str(e)}")
        
        await db.commit()
        
        return {
            "imported_count": imported_count,
            "errors": errors,
        }
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")


@router.get("/export/books")
async def export_books(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Book))
    books = result.scalars().all()
    
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Books"
    
    headers = ["ISBN", "书名", "作者", "出版社", "分类", "总数量", "可借数量", "位置"]
    sheet.append(headers)
    
    for book in books:
        sheet.append([
            book.isbn,
            book.title,
            book.author,
            book.publisher or "",
            book.category or "",
            book.total_quantity,
            book.available_quantity,
            book.location or "",
        ])
    
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename=books_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        }
    )


@router.post("/import/readers")
async def import_readers(file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Only Excel files are supported")
    
    try:
        contents = await file.read()
        workbook = load_workbook(filename=BytesIO(contents))
        sheet = workbook.active
        
        imported_count = 0
        errors = []
        
        headers = [cell.value for cell in sheet[1]]
        header_map = {}
        for idx, header in enumerate(headers):
            if header:
                header_map[str(header).strip().lower()] = idx
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                name_idx = header_map.get("name") or header_map.get("姓名")
                card_idx = header_map.get("card") or header_map.get("卡号")
                
                if name_idx is None or card_idx is None:
                    continue
                
                name = str(row[name_idx]) if row[name_idx] else None
                card = str(row[card_idx]) if row[card_idx] else None
                
                if not name or not card:
                    continue
                
                existing = await db.execute(select(Reader).where(Reader.card_number == card))
                if existing.scalar_one_or_none():
                    continue
                
                phone_idx = header_map.get("phone") or header_map.get("电话")
                email_idx = header_map.get("email") or header_map.get("邮箱")
                dept_idx = header_map.get("department") or header_map.get("部门")
                
                reader = Reader(
                    name=name,
                    card_number=card,
                    phone=str(row[phone_idx]) if phone_idx and row[phone_idx] else None,
                    email=str(row[email_idx]) if email_idx and row[email_idx] else None,
                    department=str(row[dept_idx]) if dept_idx and row[dept_idx] else None,
                )
                
                db.add(reader)
                imported_count += 1
                
            except Exception as e:
                errors.append(f"Row {row_idx}: {str(e)}")
        
        await db.commit()
        
        return {
            "imported_count": imported_count,
            "errors": errors,
        }
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Import failed: {str(e)}")


@router.get("/export/readers")
async def export_readers(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Reader))
    readers = result.scalars().all()
    
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Readers"
    
    headers = ["姓名", "卡号", "电话", "邮箱", "部门", "状态"]
    sheet.append(headers)
    
    for reader in readers:
        sheet.append([
            reader.name,
            reader.card_number,
            reader.phone or "",
            reader.email or "",
            reader.department or "",
            reader.status,
        ])
    
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename=readers_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        }
    )
